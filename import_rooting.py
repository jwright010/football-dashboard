"""
Import a week of rooting interests.

Accepts either the spreadsheet or a plain text file:

    python import_rooting.py rooting_interests.xlsx
    python import_rooting.py rooting_week_2.txt

The spreadsheet has dropdowns for Team, Stat and Direction, so the only
free entry is the player name and the number. Its rows are turned into
the same text lines the parser below reads, which means both routes go
through one set of rules and cannot drift apart.

File format, one goal per line:

    Week 2

    Jahmyr Gibbs (DET) rush yds over 85
    Amon-Ra St. Brown (DET) rec over 6
    Josh Allen (BUF) pass TD over 2
    Travis Kelce (KC) any TD
    James Cook (BUF) rush yds under 60

    BUF win
    KC points over 24
    DAL points under 17

Player lines need the team in parentheses, which is how the right game
and roster get found. Blank lines, '#' comments, and a leading 'Week N'
line are ignored.

When a name matches more than one player, you are asked to choose. The
athlete id is then stored permanently, so this only happens once.

Re-importing the same week replaces it.
"""

import re
import sys
from pathlib import Path

import db
import rooting

# The seed row shipped in the template, and the row it ships in.
# It is only skipped when found untouched in that exact row, so a real
# goal that happens to match the example is never silently dropped.
EXAMPLE_ROW = ("Jahmyr Gibbs", "DET", "rush yds", "over", 85)
EXAMPLE_ROW_NUMBER = 7


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# When several tabs hold goals, the first of these that exists wins.
PREFERRED_SHEETS = ("Rooting Revised", "Rooting")

# Header names, compared as lower-case letters only, so "BOZO?" matches
# "bozo". Columns may sit anywhere in the header row. A tab is read one of
# two ways: your own wording in a Revised column plus Team, or the
# structured Player, Team, Stat, Direction, and Target columns.
REVISED_COLUMNS = ("revised", "team")
STRUCTURED_COLUMNS = ("player", "team", "stat")
OPTIONAL = ("direction", "target", "bozo")
YES = ("yes", "y")


def _header_key(value):
    return re.sub(r"[^a-z]", "", str(value or "").lower())


def _find_header(sheet, max_rows=20):
    """(row number, {name: column index}, mode) for the goals header, or None."""
    wanted = set(REVISED_COLUMNS + STRUCTURED_COLUMNS + OPTIONAL)
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1
    ):
        columns = {}
        for index, value in enumerate(row):
            key = _header_key(value)
            if key in wanted and key not in columns:
                columns[key] = index
        if all(name in columns for name in REVISED_COLUMNS):
            return row_number, columns, "revised"
        if all(name in columns for name in STRUCTURED_COLUMNS):
            return row_number, columns, "structured"
    return None


def _find_week(sheet, header_row):
    """The number beside a "Week" label above the header, else B2."""
    for row in sheet.iter_rows(min_row=1, max_row=header_row, values_only=True):
        cells = list(row)
        for index, value in enumerate(cells):
            if _header_key(value) == "week":
                for after in cells[index + 1:]:
                    if after not in (None, ""):
                        try:
                            return int(float(after))
                        except (TypeError, ValueError):
                            return None
    value = sheet["B2"].value
    try:
        return int(float(value)) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def pick_sheet(workbook):
    """(sheet, header row, columns, mode) for the tab to import."""
    ordered = [workbook[n] for n in PREFERRED_SHEETS if n in workbook.sheetnames]
    ordered += [s for s in workbook.worksheets if s not in ordered]
    for sheet in ordered:
        found = _find_header(sheet)
        if found:
            return (sheet,) + found
    raise RuntimeError(
        "No tab has a Revised and Team header, or Player, Team, and Stat headers."
    )


# ---------------------------------------------------------------------
# Your own wording, from the Revised column
# ---------------------------------------------------------------------
#   James Cook 80+ Rush Yds        player stat, at least 80
#   James Cook 60.5- Rush Yds      player stat, 60.5 or fewer
#   Davanate Adams Anytime TD      player scores a touchdown
#   Jaguars Moneyline              team wins
#   Patriots 7.5+ Spread           team wins by more than 7.5
#   Packers 26.5+ Total Pts        team scores at least 26.5
# The team always comes from the Team column, so words naming a team,
# like "Jaguars", only need to read well. Anything else is reported,
# never guessed.

MONEYLINE = re.compile(r"^(?P<subject>.+?)\s+moneyline$", re.IGNORECASE)
ANYTIME_TD = re.compile(r"^(?P<subject>.+?)\s+anytime\s+tds?$", re.IGNORECASE)
NUMBERED = re.compile(
    r"^(?P<subject>.+?)\s+(?P<number>\d+(?:\.\d+)?)\s*(?P<sign>[+-])\s*(?P<stat>\S.*)$"
)

# Spreads written the way sportsbooks print them, with the sign in front:
#   San Francisco 49ers -2.5 Spread   favored by 2.5, so wins by 3 or more
# Only half-point lines are read. A whole-number line can push, and a "+"
# line means an underdog covering, so both are reported, never guessed.
SPREAD_FAVORITE = re.compile(
    r"^(?P<subject>.+?)\s+-(?P<number>\d+(?:\.\d+)?)\s+spread$", re.IGNORECASE
)
SPREAD_UNDERDOG = re.compile(
    r"^(?P<subject>.+?)\s+\+(?P<number>\d+(?:\.\d+)?)\s+spread$", re.IGNORECASE
)

# Longer spellings of the stat words, so "Rushing + Receiving Yds" and
# "Receptions" read the same as "Rush + Rec Yds" and "Rec".
WORD_FIXES = (
    (r"\brushing\b", "rush"), (r"\breceiving\b", "rec"),
    (r"\breceptions?\b", "rec"), (r"\bcatches\b", "rec"),
    (r"\bpassing\b", "pass"), (r"\byards?\b", "yds"),
    (r"\btouchdowns?\b", "td"), (r"\btds\b", "td"),
)


def _fix_words(words):
    for pattern, replacement in WORD_FIXES:
        words = re.sub(pattern, replacement, words)
    return words


# Team stats, by the words you'd type after the number.
TEAM_WORDS = {
    "spread": "win by",
    "total pts": "points", "total points": "points", "team total": "points",
    "pts": "points", "points": "points",
}

# Common betting words the stat list doesn't already know.
STAT_WORDS = {
    "pass tds": "pass td", "passing tds": "pass td",
    "rush tds": "rush td", "rushing tds": "rush td",
    "rec tds": "rec td", "receiving tds": "rec td",
    "passing yds": "pass yds", "rushing yds": "rush yds", "receiving yds": "rec yds",
    "rush+rec yds": "rush + rec yds", "rush & rec yds": "rush + rec yds",
    "rush and rec yds": "rush + rec yds",
    "pass attempts": "attempts", "pass completions": "completions",
    "rush attempts": "carries", "rushing attempts": "carries",
    "fgs": "field goals", "fgs made": "field goals",
}


def _number_text(value):
    number = float(value)
    return str(int(number)) if number.is_integer() else str(number)


def parse_revised(text, team):
    """(line, detail) from one Revised cell, or (None, problem).

    line is the importer's own format with your wording attached. detail
    is the wording shown on cards: for a player, the part after his name,
    since the name already shows; for a team, the whole wording.
    """
    import boxscore

    wording = " ".join(str(text).split())
    if not team:
        return None, f"{wording!r}: needs a team in the Team column"

    match = MONEYLINE.match(wording)
    if match:
        return f"{team} win", wording

    match = ANYTIME_TD.match(wording)
    if match:
        name = match.group("subject").strip()
        return f"{name} ({team}) any TD", wording[len(name):].strip()

    match = SPREAD_FAVORITE.match(wording)
    if match:
        number = match.group("number")
        if float(number).is_integer():
            return None, (f"{wording!r}: a whole-number spread can push. Write "
                          f"\"{number}.5+ Spread\" for a win by {int(number) + 1} "
                          f"or more, or \"{number}+ Spread\" for a win by {number} or more")
        return f"{team} win by at least {number}", wording

    match = SPREAD_UNDERDOG.match(wording)
    if match:
        return None, (f"{wording!r}: a \"+\" spread means an underdog covering, which "
                      "isn't read. Write it as \"3.5+ Spread\" after the team's name "
                      "for a win by 4 or more")

    match = NUMBERED.match(wording)
    if not match:
        return None, (f"{wording!r}: couldn't read it. Use forms like "
                      "\"James Cook 80+ Rush Yds\", \"James Cook 60.5- Rush Yds\", "
                      "\"Jaguars Moneyline\", or \"Patriots 7.5+ Spread\"")

    name = match.group("subject").strip()
    number = _number_text(match.group("number"))
    direction = "at least" if match.group("sign") == "+" else "under"
    words = " ".join(match.group("stat").lower().split())

    if words in TEAM_WORDS:
        team_stat = TEAM_WORDS[words]
        if team_stat == "win by" and direction == "under":
            return None, (f"{wording!r}: a spread needs \"+\", "
                          "as in \"Panthers 3.5+ Spread\"")
        return f"{team} {team_stat} {direction} {number}", wording

    words = _fix_words(words)
    phrase = STAT_WORDS.get(words, words)
    if boxscore.resolve_stat(phrase) is None:
        return None, (f"{wording!r}: unknown stat {match.group('stat')!r}. "
                      "Name it fully, such as Rush Yds, Rec Yds, or Pass Yds")
    return (f"{name} ({team}) {phrase} {direction} {number}",
            wording[len(name):].strip())


def read_workbook(path, info=None):
    """Turn the rooting tab into text lines. Returns (week, lines, notes).

    Columns are found by their header names, so added or removed columns
    shift nothing. Pass a dict as info to learn which tab was read.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError(
            "Reading .xlsx needs openpyxl. Run: pip install openpyxl"
        )

    workbook = load_workbook(path, data_only=True)
    sheet, header_row, columns, mode = pick_sheet(workbook)
    if info is not None:
        info["sheet"] = sheet.title
        info["mode"] = mode
    week = _find_week(sheet, header_row)
    example_row = header_row + 1
    has_bozo = "bozo" in columns

    def cell(row, name):
        index = columns.get(name)
        return _clean(row[index]) if index is not None and index < len(row) else ""

    lines, notes = [], []
    if mode == "revised" and not has_bozo:
        notes.append(f"The \"{sheet.title}\" tab has no BOZO? column, so nothing "
                     "shows on the BOZO tab. Add the column and mark goals Yes or Y.")

    for offset, row in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True)
    ):
        row_number = header_row + 1 + offset
        bozo = has_bozo and cell(row, "bozo").lower() in YES
        tags = " [BOZO]" if bozo else ""

        if mode == "revised":
            text, team = cell(row, "revised"), cell(row, "team")
            if not text:
                continue
            line, detail = parse_revised(text, team)
            if line is None:
                notes.append(f"Skipped {detail}")
                continue
            safe = detail.replace("{", "(").replace("}", ")")
            lines.append(f"{line} {{{safe}}}{tags}")
            continue

        player, team, stat = cell(row, "player"), cell(row, "team"), cell(row, "stat")
        direction, target = cell(row, "direction"), cell(row, "target")

        if not team and not stat:
            continue

        if row_number == example_row and (
            player, team, stat, direction, target
        ) == tuple(_clean(v) for v in EXAMPLE_ROW):
            notes.append(
                "Skipped the untouched example row. Overwrite or delete it."
            )
            continue

        if not team:
            notes.append(f"Skipped a row with no team: {player or stat!r}")
            continue
        if not stat:
            notes.append(f"Skipped {player or team!r}: no stat chosen")
            continue

        if direction.lower() in ("(none)", "none"):
            direction = ""

        # Goals that carry no number.
        if stat.lower() in ("win", "any td"):
            goal = stat
        elif direction and target:
            goal = f"{stat} {direction} {target}"
        else:
            notes.append(
                f"Skipped {player or team!r} {stat!r}: "
                "needs both a direction and a target"
            )
            continue

        line = f"{player} ({team}) {goal}" if player else f"{team} {goal}"
        lines.append(line + tags)

    return week, lines, notes


def ask(query, candidates):
    """Prompt for the right player when a name is ambiguous."""
    import roster

    print(f"\n  '{query}' matches {len(candidates)} players:")
    for index, athlete in enumerate(candidates, start=1):
        print(f"    {index}. {roster.describe(athlete)}")
    print("    0. skip this one")

    while True:
        answer = input("  Which one? ").strip()
        if answer in ("0", "s", "skip", ""):
            return None
        if answer.isdigit() and 1 <= int(answer) <= len(candidates):
            return candidates[int(answer) - 1]
        print("  Enter a number from the list.")


def infer_week(text, path):
    """Week number from a 'Week N' line, or from the filename."""
    match = re.search(r"^\s*week\s+(\d+)", text, re.IGNORECASE | re.MULTILINE)
    if match:
        return int(match.group(1))
    match = re.search(r"week[_-]?(\d+)", path, re.IGNORECASE)
    return int(match.group(1)) if match else None


def current_season():
    conn = db.connect()
    try:
        row = conn.execute("SELECT MAX(season) AS s FROM game").fetchone()
        return row["s"] if row and row["s"] else None
    finally:
        conn.close()


def main():
    if len(sys.argv) < 2:
        print("Usage: python import_rooting.py <file.txt>")
        return 1

    path = sys.argv[1]
    db.init()

    if not Path(path).exists():
        print(f"Not found: {path}")
        return 1

    season = current_season()
    if not season:
        print("No games stored yet. Run: python poller.py")
        return 1

    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            info = {}
            week, lines, notes = read_workbook(path, info)
        except Exception as exc:
            print(f"Could not read {path}: {exc}")
            return 1
        how = "your Revised wording" if info["mode"] == "revised" else "its columns"
        print(f"Read the \"{info['sheet']}\" tab, using {how}.")
        for note in notes:
            print(f"  {note}")
        if notes:
            print()
        text = "\n".join(lines)
        if week is None:
            week = infer_week(text, path)
    else:
        with open(path, encoding="utf-8") as handle:
            text = handle.read()
        week = infer_week(text, path)

    if week is None:
        print("Could not tell which week this is.")
        print("Put the number beside \"Week\" at the top of the spreadsheet, "
              "or add a 'Week 2' line to the text file.")
        return 1

    goals, errors = rooting.parse_text(text)

    if errors:
        print(f"{len(errors)} line(s) could not be read:")
        for error in errors:
            print(f"  {error}")
        print()

    if not goals:
        print("Nothing to import.")
        return 1

    print(f"Season {season}, week {week}. Read {len(goals)} goal(s).")

    stored, problems = rooting.import_goals(season, week, goals, chooser=ask)

    print(f"\nStored {stored} goal(s)")
    if problems:
        print(f"\n{len(problems)} problem(s):")
        for problem in problems:
            print(f"  {problem}")

    if not stored:
        return 1

    detail = rooting.week_detail(season, week)
    print(f"\nWeek {week}: {detail['achieved']} of {detail['total']} "
          f"achieved so far\n")

    for item in detail["items"]:
        mark = {
            "achieved": "+", "failed": "x", "live": ">",
            "pending": ".", "noshow": "?",
        }.get(item["status"], " ")
        current = "--" if item["current"] is None else f"{item['current']:g}"
        where = f"vs {item['opponent']}" if item["opponent"] else ""
        print(f"  {mark} {item['label']:24} {item['goal_text']:26} "
              f"{current:>6}   {item['status']:9} {where}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
